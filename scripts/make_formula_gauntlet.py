"""Generate fixtures/formula_gauntlet.xlsx — a self-checking formula
compatibility corpus.

Every test row computes a formula (col C), states the Excel-expected
result (col D, sometimes itself a closed-form formula), and scores
itself (col E) with a type-strict check:

  num:  =IF(ISNUMBER(C),IF(ABS(C-D)<1E-9,"OK","FAIL"),"FAIL")
  text: =IF(ISTEXT(C),IF(EXACT(C,D),"OK","FAIL"),"FAIL")
  bool: =IF(ISLOGICAL(C),IF(C=D,"OK","FAIL"),"FAIL")
  err:  =IF(IFERROR(ERROR.TYPE(C),0)=<code>,"OK","FAIL")

The same file is valid in Excel: opened there, every check should read
OK (that's how the expected values themselves get audited). Opened in
fastsheet, any non-OK row is an IronCalc compatibility gap; the "Gaps"
sheet holds cases known to fail (functions IronCalc 0.7.1 lacks).
tests/formula_gauntlet.rs loads this file headlessly and compares the
failure set against its recorded baseline.

Run: python3 scripts/make_formula_gauntlet.py
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

OUT = Path(__file__).resolve().parent.parent / "fixtures"
OUT.mkdir(exist_ok=True)

EPS = "0.000000001"
ERROR_CODES = {
    "#NULL!": 1, "#DIV/0!": 2, "#VALUE!": 3, "#REF!": 4,
    "#NAME?": 5, "#NUM!": 6, "#N/A": 7,
}


def serial(y, m, d):
    """Excel 1900-system date serial."""
    return (date(y, m, d) - date(1899, 12, 30)).days


# ---------------------------------------------------------------------------
# Case tables. Tuple: (formula, expected, kind, note)
#   kind "num" / "text" / "bool"  — expected is a literal (or a formula
#   string starting with '=' for closed-form expected values)
#   kind "err" — expected is the error name from ERROR_CODES
# ---------------------------------------------------------------------------

OPERATORS = [
    ("=2+3*4", 14, "num", "precedence: * before +"),
    ("=(2+3)*4", 20, "num", "parens"),
    ("=2^3^2", 64, "num", "^ is LEFT-associative in Excel: (2^3)^2"),
    ("=-2^2", 4, "num", "unary minus binds tighter than ^: (-2)^2"),
    ("=2^-2", 0.25, "num", "negative exponent"),
    ("=50%", 0.5, "num", "percent literal"),
    ("=50%+50%", 1, "num", "percent arithmetic"),
    ("=10%*10", 1, "num", ""),
    ("=5/2", 2.5, "num", "true division"),
    ("=0.1+0.2", 0.3, "num", "fp addition within epsilon"),
    ("=3--2", 5, "num", "double negative"),
    ("=--\"7\"", 7, "num", "double unary coerces text to number"),
    ("=\"5\"+2", 7, "num", "text-number coercion in +"),
    ("=TRUE+1", 2, "num", "boolean coerces to 1"),
    ("=TRUE*5", 5, "num", ""),
    ("=(1=1)+(2=2)", 2, "num", "comparisons coerce to 1"),
    ("=2*-3", -6, "num", "unary minus after *"),
    ("=1E3+1", 1001, "num", "scientific literal"),
    ("=7&8", "78", "text", "& concatenates numbers as text"),
    ("=\"a\"&\"b\"", "ab", "text", ""),
    ("=1&\"a\"", "1a", "text", ""),
    ("=\"5\"&2", "52", "text", "& never adds"),
    ("=1=1", True, "bool", ""),
    ("=1<>2", True, "bool", ""),
    ("=\"a\"=\"A\"", True, "bool", "= is case-insensitive for text"),
    ("=\"a\"<\"b\"", True, "bool", "text ordering"),
    ("=2>10", False, "bool", ""),
    ("=99999<\"a\"", True, "bool", "any text sorts above any number"),
    ("=TRUE>100", True, "bool", "logicals sort above numbers"),
    ("=1/0", "#DIV/0!", "err", ""),
    ("=0/0", "#DIV/0!", "err", ""),
    ("=\"a\"+1", "#VALUE!", "err", "non-numeric text in arithmetic"),
    ("=1/0+1", "#DIV/0!", "err", "errors propagate"),
]

MATH = [
    ("=ABS(-5)", 5, "num", ""),
    ("=INT(-8.5)", -9, "num", "INT floors toward -inf"),
    ("=TRUNC(-8.9)", -8, "num", "TRUNC cuts toward zero (contrast INT)"),
    ("=TRUNC(3.14159,2)", 3.14, "num", ""),
    ("=ROUND(2.5,0)", 3, "num", "half away from zero"),
    ("=ROUND(-2.5,0)", -3, "num", "half away from zero (negative)"),
    ("=ROUND(2.4,0)", 2, "num", ""),
    ("=ROUND(123.456,2)", 123.46, "num", ""),
    ("=ROUND(15,-1)", 20, "num", "negative digits round to tens"),
    ("=ROUNDUP(2.01,0)", 3, "num", ""),
    ("=ROUNDDOWN(2.99,0)", 2, "num", ""),
    ("=MOD(10,3)", 1, "num", ""),
    ("=MOD(-1,3)", 2, "num", "result takes sign of divisor"),
    ("=MOD(1,-3)", -2, "num", "result takes sign of divisor"),
    ("=MOD(1,0)", "#DIV/0!", "err", ""),
    ("=QUOTIENT(10,3)", 3, "num", ""),
    ("=CEILING(2.1,1)", 3, "num", ""),
    ("=CEILING.MATH(-2.5)", -2, "num", "default rounds toward +inf"),
    ("=FLOOR.MATH(-2.5)", -3, "num", "rounds toward -inf"),
    ("=MROUND(7,3)", 6, "num", ""),
    ("=MROUND(8,3)", 9, "num", ""),
    ("=EVEN(3)", 4, "num", ""),
    ("=EVEN(-3)", -4, "num", "away from zero"),
    ("=ODD(4)", 5, "num", ""),
    ("=ODD(-4)", -5, "num", "away from zero"),
    ("=SIGN(-3)", -1, "num", ""),
    ("=SQRT(16)", 4, "num", ""),
    ("=SQRT(-4)", "#NUM!", "err", ""),
    ("=POWER(2,10)", 1024, "num", ""),
    ("=EXP(0)", 1, "num", ""),
    ("=LN(EXP(2))", 2, "num", ""),
    ("=LOG(8,2)", 3, "num", ""),
    ("=LOG(-1)", "#NUM!", "err", ""),
    ("=LOG10(1000)", 3, "num", ""),
    ("=PI()", 3.141592653589793, "num", ""),
    ("=SIN(0)", 0, "num", ""),
    ("=COS(0)", 1, "num", ""),
    ("=SIN(PI()/2)", 1, "num", ""),
    ("=TAN(PI()/4)", 1, "num", ""),
    ("=DEGREES(PI())", 180, "num", ""),
    ("=RADIANS(180)", 3.141592653589793, "num", ""),
    ("=FACT(5)", 120, "num", ""),
    ("=FACTDOUBLE(6)", 48, "num", ""),
    ("=COMBIN(5,2)", 10, "num", ""),
    ("=GCD(12,18)", 6, "num", ""),
    ("=LCM(4,6)", 12, "num", ""),
    ("=SUMSQ(3,4)", 25, "num", ""),
    ("=PRODUCT(2,3,4)", 24, "num", ""),
    ("=SUM(Data!E1:E8)", 39, "num", "range sum"),
    ("=SUM({1,2,3,4})", 10, "num", "array constant"),
    ("=SUMIF(Data!E1:E8,\">4\")", 29, "num", ""),
    ("=SUMIFS(Data!B2:B6,Data!A2:A6,\"Banana\")", 20, "num", ""),
    ("=SUBTOTAL(9,Data!E1:E8)", 39, "num", ""),
    ("=SUMPRODUCT(Data!B2:B6,Data!C2:C6)", 149.5, "num", ""),
    ("=SUMPRODUCT(Data!E1:E8)", 39, "num", "single array just sums"),
    ("=SUMPRODUCT(Data!A2:A6,Data!B2:B6)", 0, "num", "text elements count as 0"),
    ("=SUMPRODUCT(Data!B2:B6,Data!C2:C5)", "#VALUE!", "err", "dimension mismatch"),
    ("=ROMAN(1990)", "MCMXC", "text", ""),
    ("=ARABIC(\"XIV\")", 14, "num", ""),
    ("=BASE(255,16)", "FF", "text", ""),
    ("=DECIMAL(\"FF\",16)", 255, "num", ""),
]

TEXT = [
    ("=LEN(\"hello\")", 5, "num", ""),
    ("=LEN(\"\")", 0, "num", ""),
    ("=LEFT(\"spreadsheet\",6)", "spread", "text", ""),
    ("=RIGHT(\"spreadsheet\",5)", "sheet", "text", ""),
    ("=MID(\"spreadsheet\",7,5)", "sheet", "text", ""),
    ("=LOWER(\"MiXeD\")", "mixed", "text", ""),
    ("=UPPER(\"mixed\")", "MIXED", "text", ""),
    ("=TRIM(\"  a  b  \")", "a b", "text", "collapses inner runs too"),
    ("=LEN(TRIM(Data!H2))", 6, "num", "trims '  padded  '"),
    ("=REPT(\"ab\",3)", "ababab", "text", ""),
    ("=EXACT(\"a\",\"A\")", False, "bool", "EXACT is case-sensitive"),
    ("=CONCAT(\"a\",\"b\",\"c\")", "abc", "text", ""),
    ("=CONCATENATE(\"x\",1)", "x1", "text", ""),
    ("=FIND(\"b\",\"abcabc\")", 2, "num", ""),
    ("=FIND(\"b\",\"abcabc\",3)", 5, "num", "start position"),
    ("=FIND(\"B\",\"abc\")", "#VALUE!", "err", "FIND is case-sensitive"),
    ("=SEARCH(\"B\",\"abc\")", 2, "num", "SEARCH is case-insensitive"),
    ("=SUBSTITUTE(\"aaa\",\"a\",\"b\")", "bbb", "text", ""),
    ("=SUBSTITUTE(\"aaa\",\"a\",\"b\",2)", "aba", "text", "nth instance only"),
    ("=T(\"x\")", "x", "text", ""),
    ("=VALUE(\"42\")", 42, "num", ""),
    ("=VALUE(\"3.14\")", 3.14, "num", ""),
    ("=TEXT(1234.5678,\"0.00\")", "1234.57", "text", ""),
    ("=TEXT(0.25,\"0%\")", "25%", "text", ""),
    ("=TEXT(DATE(2026,8,3),\"yyyy-mm-dd\")", "2026-08-03", "text", ""),
    ("=CHAR(65)", "A", "text", ""),
    ("=UNICODE(\"A\")", 65, "num", ""),
    ("=TEXTJOIN(\"-\",TRUE,\"a\",\"b\",\"c\")", "a-b-c", "text", ""),
    ("=TEXTJOIN(\",\",TRUE,\"a\",\"\",\"b\")", "a,b", "text", "skip empties"),
    ("=TEXTBEFORE(\"a-b-c\",\"-\")", "a", "text", ""),
    ("=TEXTBEFORE(\"a-b-c\",\"-\",2)", "a-b", "text", ""),
    ("=TEXTAFTER(\"a-b-c\",\"-\")", "b-c", "text", ""),
    # expected is written as a formula producing the text "=1+1" — a plain
    # string starting with "=" would be stored as a formula by openpyxl.
    ("=FORMULATEXT(Data!J1)", '="="&"1+1"', "text", ""),
    ("=PROPER(\"hello world\")", "Hello World", "text", ""),
    ("=PROPER(\"2-cent's worth\")", "2-Cent'S Worth", "text", "any non-letter is a word break"),
    ("=PROPER(Data!H3)", "Mixed Case", "text", ""),
    ("=CODE(\"A\")", 65, "num", ""),
    ("=CODE(\"€\")", 128, "num", "Win-1252 C1 range"),
    ("=CHAR(CODE(\"a\"))", "a", "text", "round-trip"),
    ("=CODE(\"\")", "#VALUE!", "err", "empty text"),
    ("=REPLACE(\"abcdef\",2,3,\"XY\")", "aXYef", "text", ""),
    ("=REPLACE(\"abc\",9,2,\"Z\")", "abcZ", "text", "start past end appends"),
    ("=REPLACE(\"abc\",1,0,\"Z\")", "Zabc", "text", "zero chars inserts"),
    ("=REPLACE(\"abc\",0,1,\"Z\")", "#VALUE!", "err", "start below 1"),
    ("=FIXED(1234.567,1)", "1,234.6", "text", ""),
    ("=FIXED(1234.567)", "1,234.57", "text", "default 2 decimals"),
    ("=FIXED(-1234.567,1)", "-1,234.6", "text", ""),
    ("=FIXED(1234.567,-2)", "1,200", "text", "negative decimals round left"),
    ("=FIXED(1234.567,1,TRUE)", "1234.6", "text", "no_commas"),
    ("=NUMBERVALUE(\"1.5\")", 1.5, "num", ""),
    ("=NUMBERVALUE(\"2.5%\")", 0.025, "num", "trailing percent divides"),
    ("=NUMBERVALUE(\"3,50\",\",\",\".\")", 3.5, "num", "custom separators"),
    ("=NUMBERVALUE(\"1,234.5\")", 1234.5, "num", "group separators stripped"),
    ("=NUMBERVALUE(\"\")", 0, "num", "empty text is zero"),
    ("=NUMBERVALUE(\"abc\")", "#VALUE!", "err", ""),
]

LOGICAL = [
    ("=IF(1>2,\"y\",\"n\")", "n", "text", ""),
    ("=IF(0,\"y\",\"n\")", "n", "text", "0 is falsy"),
    ("=IF(3,\"y\",\"n\")", "y", "text", "nonzero is truthy"),
    ("=AND(TRUE,1=1)", True, "bool", ""),
    ("=AND(TRUE,FALSE)", False, "bool", ""),
    ("=OR(FALSE,FALSE)", False, "bool", ""),
    ("=OR(0,1)", True, "bool", "numeric coercion"),
    ("=XOR(TRUE,TRUE)", False, "bool", ""),
    ("=XOR(TRUE,FALSE,TRUE)", False, "bool", "even count of TRUE"),
    ("=NOT(FALSE)", True, "bool", ""),
    ("=IFERROR(1/0,\"caught\")", "caught", "text", ""),
    ("=IFERROR(42,\"x\")", 42, "num", "no error passes through"),
    ("=IFNA(NA(),\"na\")", "na", "text", ""),
    ("=IFNA(1/0,\"na\")", "#DIV/0!", "err", "IFNA only catches #N/A"),
    ("=IFS(FALSE,\"a\",TRUE,\"b\")", "b", "text", ""),
    ("=IFS(FALSE,\"a\")", "#N/A", "err", "no branch matched"),
    ("=SWITCH(2,1,\"one\",2,\"two\",\"def\")", "two", "text", ""),
    ("=SWITCH(9,1,\"one\",\"def\")", "def", "text", "default arm"),
    ("=TRUE()", True, "bool", ""),
]

LOOKUP = [
    ("=VLOOKUP(\"Cherry\",Data!A2:C6,3,FALSE)", 3.2, "num", ""),
    ("=VLOOKUP(\"cherry\",Data!A2:C6,3,FALSE)", 3.2, "num", "case-insensitive"),
    ("=VLOOKUP(\"Zzz\",Data!A2:C6,2,FALSE)", "#N/A", "err", ""),
    ("=VLOOKUP(\"Coconut\",Data!A2:C6,2,TRUE)", 5, "num", "approx falls back to Cherry"),
    ("=HLOOKUP(\"b\",Data!K1:M2,2,FALSE)", 2, "num", ""),
    ("=INDEX(Data!A2:C6,2,1)", "Banana", "text", ""),
    ("=INDEX(Data!E1:E8,4)", 1, "num", "vector form"),
    ("=MATCH(8,Data!E1:E8,0)", 3, "num", ""),
    ("=MATCH(\"Date\",Data!A2:A6,0)", 4, "num", ""),
    ("=MATCH(35,Data!F3:F7,1)", 3, "num", "approx over sorted 10..50"),
    ("=INDEX(Data!C2:C6,MATCH(\"Banana\",Data!A2:A6,0))", 0.25, "num", "INDEX+MATCH"),
    ("=CHOOSE(3,\"a\",\"b\",\"c\")", "c", "text", ""),
    ("=OFFSET(Data!A2,1,1)", 20, "num", "lands on B3"),
    ("=SUM(OFFSET(Data!E1,0,0,4,1))", 17, "num", "resized range 5+3+8+1"),
    ("=ROWS(Data!A2:C6)", 5, "num", ""),
    ("=COLUMNS(Data!A2:C6)", 3, "num", ""),
    ("=INDIRECT(\"Data!B3\")", 20, "num", ""),
    ("=INDIRECT(\"Data!R3C2\",FALSE)", 20, "num", "R1C1 mode"),
    ("=SUM(INDIRECT(\"Data!R1C5:R4C5\",FALSE))", 17, "num", "R1C1 range"),
    ("=INDIRECT(\"TaxRate\")", 0.15, "num", "defined name through INDIRECT"),
    ("=VLOOKUP(\"Date\",PriceTable,2,FALSE)", 15, "num", "defined-name table"),
    ("=TaxRate*100", 15, "num", "defined name in arithmetic"),
    ("=XLOOKUP(\"Date\",Data!A2:A6,Data!B2:B6)", 15, "num", ""),
    ("=XLOOKUP(\"Zzz\",Data!A2:A6,Data!B2:B6,\"missing\")", "missing", "text", ""),
    ("=LOOKUP(35,Data!F3:F7)", 30, "num", "vector LOOKUP approx"),
    ("=ERROR.TYPE(1/0)", 2, "num", ""),
    ("=TYPE(1)", 1, "num", ""),
    ("=TYPE(\"a\")", 2, "num", ""),
    ("=TYPE(TRUE)", 4, "num", ""),
    ("=N(TRUE)", 1, "num", ""),
    ("=N(\"a\")", 0, "num", ""),
    ("=ADDRESS(1,1)", "$A$1", "text", ""),
    ("=ADDRESS(2,3,2)", "C$2", "text", "row absolute"),
    ("=ADDRESS(2,3,3)", "$C2", "text", "column absolute"),
    ("=ADDRESS(2,3,4)", "C2", "text", "relative"),
    ("=ADDRESS(2,3,1,FALSE)", "R2C3", "text", "R1C1 absolute"),
    ("=ADDRESS(2,3,4,FALSE)", "R[2]C[3]", "text", "R1C1 relative"),
    ("=ADDRESS(1,1,1,TRUE,\"Data\")", "Data!$A$1", "text", "sheet prefix"),
    ("=ADDRESS(1,1,1,TRUE,\"My Sheet\")", "'My Sheet'!$A$1", "text", "quoted sheet"),
    ("=ADDRESS(1,16385)", "#VALUE!", "err", "beyond last column"),
]

DATES = [
    ("=DATE(2026,8,3)", serial(2026, 8, 3), "num", ""),
    ("=YEAR(DATE(2026,8,3))", 2026, "num", ""),
    ("=MONTH(DATE(2026,8,3))", 8, "num", ""),
    ("=DAY(DATE(2026,8,3))", 3, "num", ""),
    ("=DATE(2026,13,1)", serial(2027, 1, 1), "num", "month overflow rolls"),
    ("=DATE(2026,2,30)", serial(2026, 3, 2), "num", "day overflow rolls"),
    ("=WEEKDAY(DATE(2026,8,3))", 2, "num", "2026-08-03 is a Monday"),
    ("=WEEKDAY(DATE(2026,8,3),2)", 1, "num", "type 2: Monday=1"),
    ("=EDATE(DATE(2026,1,31),1)", serial(2026, 2, 28), "num", "clamps to month end"),
    ("=EOMONTH(DATE(2026,2,10),0)", serial(2026, 2, 28), "num", ""),
    ("=EOMONTH(DATE(2024,2,10),0)", serial(2024, 2, 29), "num", "leap year"),
    ("=DAYS(DATE(2026,12,31),DATE(2026,1,1))", 364, "num", ""),
    ("=DAYS360(DATE(2026,1,1),DATE(2026,12,31))", 360, "num", "US method"),
    ("=DATEDIF(DATE(2020,1,15),DATE(2026,8,3),\"Y\")", 6, "num", ""),
    ("=DATEDIF(DATE(2020,1,15),DATE(2026,8,3),\"M\")", 78, "num", ""),
    ("=TIME(12,30,0)", 0.5208333333333334, "num", ""),
    ("=HOUR(0.75)", 18, "num", ""),
    ("=MINUTE(TIME(1,23,45))", 23, "num", ""),
    ("=SECOND(TIME(1,2,3))", 3, "num", ""),
    ("=DATEVALUE(\"2026-08-03\")", serial(2026, 8, 3), "num", "ISO parse"),
    ("=TIMEVALUE(\"12:00\")", 0.5, "num", ""),
    ("=WEEKNUM(DATE(2026,1,1))", 1, "num", ""),
    ("=ISOWEEKNUM(DATE(2026,1,1))", 1, "num", "Jan 1 2026 is a Thursday"),
    ("=NETWORKDAYS(DATE(2026,8,3),DATE(2026,8,7))", 5, "num", "Mon..Fri"),
    ("=WORKDAY(DATE(2026,8,7),1)", serial(2026, 8, 10), "num", "Fri+1 skips weekend"),
    ("=YEARFRAC(DATE(2026,1,1),DATE(2027,1,1))", 1, "num", "basis 0"),
    ("=YEAR(TODAY())>=2026", True, "bool", "shape check only"),
    ("=NOW()>=TODAY()", True, "bool", "shape check only"),
]

STATS = [
    ("=AVERAGE(Data!E1:E8)", 4.875, "num", ""),
    ("=AVERAGE(1,2,3)", 2, "num", ""),
    ("=AVERAGEA(1,TRUE)", 1, "num", "TRUE counts as 1"),
    ("=MEDIAN(Data!E1:E8)", 4.5, "num", ""),
    ("=MAX(Data!E1:E8)", 9, "num", ""),
    ("=MIN(Data!E1:E8)", 1, "num", ""),
    ("=LARGE(Data!E1:E8,2)", 8, "num", ""),
    ("=SMALL(Data!E1:E8,3)", 3, "num", ""),
    ("=COUNT(Data!G1:G5)", 1, "num", "only the number; TRUE/text/error don't count"),
    ("=COUNTA(Data!G1:G5)", 4, "num", "everything but the blank"),
    ("=COUNTBLANK(Data!G1:G5)", 1, "num", ""),
    ("=COUNTIF(Data!E1:E8,\">4\")", 4, "num", ""),
    ("=COUNTIF(Data!A2:A6,\"*e*\")", 4, "num", "wildcard contains"),
    ("=COUNTIFS(Data!E1:E8,\">2\",Data!E1:E8,\"<8\")", 4, "num", ""),
    ("=AVERAGEIF(Data!E1:E8,\">4\")", 7.25, "num", ""),
    ("=STDEV.P(2,4,4,4,5,5,7,9)", 2, "num", "exact"),
    ("=VAR.P(2,4,4,4,5,5,7,9)", 4, "num", "exact"),
    ("=RANK.EQ(7,Data!E1:E8)", 3, "num", "descending default"),
    ("=GEOMEAN(2,8)", 4, "num", ""),
    ("=HARMEAN(2,6)", 3, "num", ""),
    ("=DEVSQ(1,2,3)", 2, "num", ""),
    ("=MAXIFS(Data!B2:B6,Data!A2:A6,\"<>Banana\")", 15, "num", ""),
    ("=MINIFS(Data!B2:B6,Data!A2:A6,\"<>Banana\")", 5, "num", ""),
    ("=MODE(1,2,2,3)", 2, "num", ""),
    ("=MODE(3,3,2,2)", 3, "num", "tie goes to first seen"),
    ("=MODE(1,2,3)", "#N/A", "err", "no repeats"),
    ("=MODE(Data!E1:E8)", "#N/A", "err", "all distinct"),
]

INFO = [
    ("=ISBLANK(Data!G4)", True, "bool", ""),
    ("=ISBLANK(\"\")", False, "bool", "empty text is not blank"),
    ("=ISNUMBER(1)", True, "bool", ""),
    ("=ISNUMBER(\"1\")", False, "bool", "no coercion in IS*"),
    ("=ISTEXT(\"a\")", True, "bool", ""),
    ("=ISTEXT(1)", False, "bool", ""),
    ("=ISLOGICAL(TRUE)", True, "bool", ""),
    ("=ISLOGICAL(1)", False, "bool", ""),
    ("=ISNA(NA())", True, "bool", ""),
    ("=ISNA(1/0)", False, "bool", ""),
    ("=ISERR(1/0)", True, "bool", "ISERR = all errors except #N/A"),
    ("=ISERR(NA())", False, "bool", ""),
    ("=ISERROR(NA())", True, "bool", ""),
    ("=ISEVEN(4)", True, "bool", ""),
    ("=ISEVEN(2.5)", True, "bool", "truncates first"),
    ("=ISODD(4)", False, "bool", ""),
    ("=ISNONTEXT(1)", True, "bool", ""),
    ("=ISFORMULA(Data!J1)", True, "bool", ""),
    ("=ISREF(Data!A1)", True, "bool", ""),
]

FINANCIAL = [
    ("=PMT(0.05/12,360,200000)",
     "=-200000*(0.05/12)/(1-(1+0.05/12)^-360)", "num", "vs closed form"),
    ("=FV(0.1,2,0,-100)", 121, "num", ""),
    ("=PV(0.1,1,0,-110)", 100, "num", ""),
    ("=NPER(0.1,0,-100,121)", 2, "num", ""),
    ("=RATE(2,0,-100,121)", 0.1, "num", ""),
    ("=NPV(0.1,100,200)", "=100/1.1+200/1.1^2", "num", "vs closed form"),
    ("=IRR(Data!O1:O2)", 0.1, "num", "-100 then 110"),
    ("=SLN(10000,1000,5)", 1800, "num", ""),
    ("=DDB(10000,1000,5,1)", 4000, "num", ""),
    ("=SYD(10000,1000,5,1)", 3000, "num", ""),
    ("=EFFECT(0.12,12)", "=1.01^12-1", "num", "vs closed form"),
    ("=NOMINAL(0.126825030131969,12)", 0.12, "num", ""),
]

ENGINEERING = [
    ("=BIN2DEC(1010)", 10, "num", ""),
    ("=DEC2BIN(10)", "1010", "text", ""),
    ("=HEX2DEC(\"FF\")", 255, "num", ""),
    ("=DEC2HEX(255)", "FF", "text", ""),
    ("=OCT2DEC(17)", 15, "num", ""),
    ("=DEC2OCT(15)", "17", "text", ""),
    ("=BITAND(12,10)", 8, "num", ""),
    ("=BITOR(12,10)", 14, "num", ""),
    ("=BITXOR(12,10)", 6, "num", ""),
    ("=BITLSHIFT(1,4)", 16, "num", ""),
    ("=BITRSHIFT(16,2)", 4, "num", ""),
    ("=DELTA(5,5)", 1, "num", ""),
    ("=DELTA(5,4)", 0, "num", ""),
    ("=GESTEP(5,4)", 1, "num", ""),
    ("=ERF(0)", 0, "num", ""),
    ("=COMPLEX(3,4)", "3+4i", "text", ""),
    ("=IMABS(\"3+4i\")", 5, "num", ""),
    ("=IMREAL(\"3+4i\")", 3, "num", ""),
    ("=IMAGINARY(\"3+4i\")", 4, "num", ""),
    ("=IMSUM(\"1+2i\",\"3+4i\")", "4+6i", "text", ""),
    ("=CONVERT(1,\"in\",\"cm\")", 2.54, "num", ""),
    ("=CONVERT(100,\"C\",\"F\")", 212, "num", ""),
]

DATABASE = [
    ("=DSUM(Data!A1:C6,\"Qty\",Data!K4:K5)", 45, "num", "Qty > 8"),
    ("=DCOUNT(Data!A1:C6,\"Qty\",Data!K4:K5)", 3, "num", ""),
    ("=DMAX(Data!A1:C6,\"Qty\",Data!K4:K5)", 20, "num", ""),
    ("=DMIN(Data!A1:C6,\"Qty\",Data!K4:K5)", 10, "num", ""),
    ("=DGET(Data!A1:C6,\"Price\",Data!L4:L5)", 4.5, "num", "Fruit = Date"),
]

ERRORS = [
    ("=OFFSET(Data!A1,-1,0)", "#REF!", "err", "walks off the grid"),
    ("=INDIRECT(\"bogus ref\")", "#REF!", "err", ""),
    ("=VLOOKUP(\"zzz\",Data!A2:C6,2,FALSE)", "#N/A", "err", ""),
    ("=NA()", "#N/A", "err", ""),
    ("=DATE(-1,1,1)", "#NUM!", "err", ""),
    ("=NOSUCHFUNCTION(1)", "#NAME?", "err", "unknown function"),
]

# Cases that pass in Excel but are known IronCalc gaps (missing
# functions). They live on their own sheet so the Summary separates
# genuine regressions from documented gaps. SUMPRODUCT / ADDRESS /
# PROPER / CODE / MODE / REPLACE / NUMBERVALUE / FIXED graduated to
# their category sheets when fastsheet implemented them (2026-08-03).
GAPS = [
    ("=SUM(TRANSPOSE(Data!E1:E8))", 39, "num", "no native TRANSPOSE (MY* only)"),
    ("=SUM(SEQUENCE(3))", 6, "num", "no SEQUENCE"),
    ("=LET(x,2,x*3)", 6, "num", "no LET"),
    ("=TEXTSPLIT(\"a,b\",\",\")", "a", "text", "no TEXTSPLIT (anchor cell)"),
]

CATEGORIES = [
    ("Operators", OPERATORS),
    ("Math", MATH),
    ("Text", TEXT),
    ("Logical", LOGICAL),
    ("Lookup", LOOKUP),
    ("Dates", DATES),
    ("Stats", STATS),
    ("Info", INFO),
    ("Financial", FINANCIAL),
    ("Engineering", ENGINEERING),
    ("Database", DATABASE),
    ("Errors", ERRORS),
    ("Gaps", GAPS),
]


def check_formula(kind, expected, row):
    c, d = f"C{row}", f"D{row}"
    if kind == "num":
        return f'=IF(ISNUMBER({c}),IF(ABS({c}-{d})<{EPS},"OK","FAIL"),"FAIL")'
    if kind == "text":
        return f'=IF(ISTEXT({c}),IF(EXACT({c},{d}),"OK","FAIL"),"FAIL")'
    if kind == "bool":
        return f'=IF(ISLOGICAL({c}),IF({c}={d},"OK","FAIL"),"FAIL")'
    if kind == "err":
        code = ERROR_CODES[expected]
        return f'=IF(IFERROR(ERROR.TYPE({c}),0)={code},"OK","FAIL")'
    raise ValueError(kind)


def build_data_sheet(ws):
    ws["A1"] = "Fruit"; ws["B1"] = "Qty"; ws["C1"] = "Price"
    fruit = [("Apple", 10, 0.5), ("Banana", 20, 0.25), ("Cherry", 5, 3.2),
             ("Date", 15, 4.5), ("Elderberry", 8, 7.0)]
    for i, (f, q, p) in enumerate(fruit, start=2):
        ws.cell(row=i, column=1, value=f)
        ws.cell(row=i, column=2, value=q)
        ws.cell(row=i, column=3, value=p)
    for i, v in enumerate([5, 3, 8, 1, 9, 2, 7, 4], start=1):   # E1:E8
        ws.cell(row=i, column=5, value=v)
    ws["F1"] = 0.15                                             # TaxRate
    for i, v in enumerate([10, 20, 30, 40, 50], start=3):       # F3:F7 sorted
        ws.cell(row=i, column=6, value=v)
    ws["G1"] = "text"; ws["G2"] = 42; ws["G3"] = True           # G4 blank
    ws["G5"] = "=NA()"
    ws["H1"] = "Hello World"; ws["H2"] = "  padded  "; ws["H3"] = "MiXeD cAsE"
    ws["J1"] = "=1+1"                                           # FORMULATEXT target
    ws["K1"] = "a"; ws["L1"] = "b"; ws["M1"] = "c"              # HLOOKUP block
    ws["K2"] = 1; ws["L2"] = 2; ws["M2"] = 3
    ws["K4"] = "Qty"; ws["K5"] = ">8"                           # DSUM criteria
    ws["L4"] = "Fruit"; ws["L5"] = "Date"                       # DGET criteria
    ws["O1"] = -100; ws["O2"] = 110                             # IRR cashflows


def main():
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    data = wb.create_sheet("Data")
    build_data_sheet(data)

    bold = Font(bold=True)
    totals = []
    for name, cases in CATEGORIES:
        ws = wb.create_sheet(name)
        for col, hdr in enumerate(["ID", "Note", "Actual", "Expected", "Check"], 1):
            cell = ws.cell(row=1, column=col, value=hdr)
            cell.font = bold
        for i, (formula, expected, kind, note) in enumerate(cases):
            row = i + 2
            ws.cell(row=row, column=1, value=f"{name[:3].upper()}{i+1:03d}")
            ws.cell(row=row, column=2, value=note or "")
            ws.cell(row=row, column=3, value=formula)
            ws.cell(row=row, column=4, value=expected)
            ws.cell(row=row, column=5, value=check_formula(kind, expected, row))
        ws.freeze_panes = "A2"
        for col, width in zip("ABCDE", (10, 42, 48, 22, 8)):
            ws.column_dimensions[col].width = width
        totals.append((name, len(cases)))

    summary["A1"] = "Category"; summary["B1"] = "Cases"
    summary["C1"] = "OK"; summary["D1"] = "Fails"
    for c in "ABCD":
        summary[f"{c}1"].font = bold
    for i, (name, count) in enumerate(totals, start=2):
        last = count + 1
        summary.cell(row=i, column=1, value=name)
        summary.cell(row=i, column=2, value=count)
        summary.cell(row=i, column=3, value=f"=COUNTIF('{name}'!$E$2:$E${last},\"OK\")")
        summary.cell(row=i, column=4, value=f"=B{i}-C{i}")
    total_row = len(totals) + 2
    summary.cell(row=total_row, column=1, value="TOTAL").font = bold
    summary.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row-1})")
    summary.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row-1})")
    summary.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})")
    summary.cell(row=total_row + 2, column=1,
                 value=f'=IF(D{total_row}=0,"ALL PASS",D{total_row}&" FAILURES")')
    for col, width in zip("ABCD", (16, 8, 8, 8)):
        summary.column_dimensions[col].width = width

    wb.defined_names["PriceTable"] = DefinedName(
        "PriceTable", attr_text="Data!$A$2:$C$6")
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Data!$F$1")

    out = OUT / "formula_gauntlet.xlsx"
    wb.save(out)
    n = sum(c for _, c in totals)
    print(f"wrote {out} — {n} cases across {len(totals)} category sheets")


if __name__ == "__main__":
    main()
