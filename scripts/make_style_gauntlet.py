"""Generate fixtures/style_gauntlet.xlsx — a self-describing style /
formatting corpus (the styling counterpart of the formula gauntlet).

Layout of every category sheet: each row is one case —
  A: ID   B: human note   C: the STYLED target cell   D: spec (JSON)

The spec is a JSON object describing what the loaded IronCalc style of
cell C must look like. tests/style_gauntlet.rs parses column D and
asserts column C's style matches, twice: once after the xlsx load, and
once more after a save-to-.xls + reload round-trip (which exercises the
BIFF style pipeline: FONT/XF/PALETTE/FORMAT emission).

Spec keys (all optional):
  font:      {"name": str, "sz": int, "b": bool, "i": bool, "u": bool,
              "strike": bool, "color": "#RRGGBB"}
  fill:      "#RRGGBB"            (solid foreground)
  numfmt:    str                  (expected style.num_fmt)
  shows:     str                  (expected get_formatted_cell_value)
  halign:    "left|center|right|justify|fill|centercontinuous|distributed"
  valign:    "top|center|bottom|justify|distributed"
  wrap:      true
  border:    {"left|right|top|bottom": {"style": str, "color": "#RRGGBB"?}}
  merge:     "C10:E12"            (must appear in the sheet's merge list)
  colwidth:  {"col": int, "chars": float}
  rowheight_pt: float             (height of the spec row itself)
  rowhidden: true                 (the spec row itself is hidden)
  colhidden: int                  (that column is hidden)
  frozen:    {"rows": int, "cols": int}
  survive:   true                 (Gaps sheet: only assert the value
                                   loads intact — feature not modeled)

Run: python3 scripts/make_style_gauntlet.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side

OUT = Path(__file__).resolve().parent.parent / "fixtures"
OUT.mkdir(exist_ok=True)

wb = Workbook()


def sheet(name):
    ws = wb.create_sheet(name)
    ws["A1"] = "ID"
    ws["B1"] = "Note"
    ws["C1"] = "Styled"
    ws["D1"] = "Spec"
    for c in "ABCD":
        ws[f"{c}1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 60
    ws.freeze_panes = "A2"
    return ws


def case(ws, row, note, value, spec, style=None, number_format=None):
    ws.cell(row=row, column=1, value=f"{ws.title[:3].upper()}{row:03d}")
    ws.cell(row=row, column=2, value=note)
    c = ws.cell(row=row, column=3, value=value)
    if style:
        for attr, obj in style.items():
            setattr(c, attr, obj)
    if number_format:
        c.number_format = number_format
    ws.cell(row=row, column=4, value=json.dumps(spec))


# ---------------------------------------------------------------------------
# Fonts — faces and sizes
# ---------------------------------------------------------------------------
ws = sheet("Fonts")
row = 2
for face in ["Arial", "Calibri", "Times New Roman", "Courier New", "Verdana",
             "Tahoma", "Segoe UI", "Georgia", "Comic Sans MS", "Consolas"]:
    case(ws, row, f"face {face}", f"The quick brown fox — {face}",
         {"font": {"name": face, "sz": 11}},
         style={"font": Font(name=face, size=11)})
    row += 1
for size in [8, 9, 10, 11, 12, 14, 16, 20, 24, 36, 48, 72]:
    case(ws, row, f"Calibri {size}pt", f"{size} point",
         {"font": {"name": "Calibri", "sz": size}},
         style={"font": Font(name="Calibri", size=size)})
    row += 1

# ---------------------------------------------------------------------------
# FontStyles — bold / italic / underline / strike and combos
# ---------------------------------------------------------------------------
ws = sheet("FontStyles")
row = 2
combos = [
    ("bold", {"b": True}, Font(bold=True)),
    ("italic", {"i": True}, Font(italic=True)),
    ("underline", {"u": True}, Font(underline="single")),
    ("double underline (collapses to single-bit u)", {"u": True},
     Font(underline="double")),
    ("strikethrough", {"strike": True}, Font(strike=True)),
    ("bold italic", {"b": True, "i": True}, Font(bold=True, italic=True)),
    ("bold underline", {"b": True, "u": True},
     Font(bold=True, underline="single")),
    ("italic strike", {"i": True, "strike": True},
     Font(italic=True, strike=True)),
    ("all four", {"b": True, "i": True, "u": True, "strike": True},
     Font(bold=True, italic=True, underline="single", strike=True)),
    ("bold 16pt Arial", {"name": "Arial", "sz": 16, "b": True},
     Font(name="Arial", size=16, bold=True)),
]
for note, fspec, fobj in combos:
    case(ws, row, note, note, {"font": fspec}, style={"font": fobj})
    row += 1

# ---------------------------------------------------------------------------
# Colors — font colors and solid fills (standard + custom RGB)
# ---------------------------------------------------------------------------
ws = sheet("Colors")
row = 2
palette = [
    ("black", "000000"), ("red", "FF0000"), ("green", "00B050"),
    ("blue", "0070C0"), ("yellow", "BF9000"), ("orange", "ED7D31"),
    ("purple", "7030A0"), ("teal", "1F9E9E"), ("gray", "808080"),
    ("custom #123456", "123456"), ("custom #AB12CD", "AB12CD"),
]
for name, rgb in palette:
    case(ws, row, f"font color {name}", f"colored text {rgb}",
         {"font": {"color": f"#{rgb}"}},
         style={"font": Font(color=f"FF{rgb}")})
    row += 1
fills = [
    ("yellow", "FFFF00"), ("light green", "C6EFCE"), ("light red", "FFC7CE"),
    ("light blue", "BDD7EE"), ("orange", "FFC000"), ("gray", "D9D9D9"),
    ("dark blue", "1F4E79"), ("black", "000000"),
    ("custom #123456", "123456"), ("custom #FEDCBA", "FEDCBA"),
]
for name, rgb in fills:
    case(ws, row, f"solid fill {name}", f"fill {rgb}",
         {"fill": f"#{rgb}"},
         style={"fill": PatternFill(fill_type="solid", fgColor=f"FF{rgb}")})
    row += 1
case(ws, row, "white on dark blue", "inverse video",
     {"font": {"color": "#FFFFFF", "b": True}, "fill": "#1F4E79"},
     style={"font": Font(color="FFFFFFFF", bold=True),
            "fill": PatternFill(fill_type="solid", fgColor="FF1F4E79")})
row += 1
case(ws, row, "red bold on yellow", "warning style",
     {"font": {"color": "#FF0000", "b": True}, "fill": "#FFFF00"},
     style={"font": Font(color="FFFF0000", bold=True),
            "fill": PatternFill(fill_type="solid", fgColor="FFFFFF00")})

# ---------------------------------------------------------------------------
# NumFmts — format string survives the load AND renders like Excel
# ---------------------------------------------------------------------------
ws = sheet("NumFmts")
row = 2
D = 46238  # 2026-08-04 (a Tuesday)
T = 49058 / 86400  # 13:37:38
cases = [
    ("integer", 1234.5678, "0", "1235"),
    ("two decimals", 1234.5678, "0.00", "1234.57"),
    ("thousands", 1234567, "#,##0", "1,234,567"),
    ("thousands + decimals", 1234.5, "#,##0.00", "1,234.50"),
    ("percent", 0.25, "0%", "25%"),
    ("percent decimals", 0.2568, "0.00%", "25.68%"),
    ("scientific", 1234.5678, "0.00E+00", "1.23E+03"),
    ("fraction", 5.25, "# ?/?", "5 1/4"),
    ("fraction 2-digit", 5.3, "# ??/??", "5 3/10"),
    ("currency", 1234.5, "$#,##0.00", "$1,234.50"),
    ("currency negative parens", -1234.5, "$#,##0.00;($#,##0.00)", "($1,234.50)"),
    ("negative red section", -5, "0.00;[Red]-0.00", "-5.00"),
    ("literal prefix", 1234, '"kr "#,##0', "kr 1,234"),
    ("literal suffix", 12, '#,##0" units"', "12 units"),
    ("zip zeros", 42, "000000", "000042"),
    ("date m/d/yyyy", D, "m/d/yyyy", "8/4/2026"),
    ("date d-mmm-yy", D, "d-mmm-yy", "4-Aug-26"),
    ("date dd/mm/yyyy", D, "dd/mm/yyyy", "04/08/2026"),
    ("date iso", D, "yyyy-mm-dd", "2026-08-04"),
    ("month-year", D, "mmm-yy", "Aug-26"),
    ("weekday name", D, "dddd", "Tuesday"),
    ("time h:mm", 0.5208333333333334, "h:mm", "12:30"),
    ("time h:mm:ss", T, "h:mm:ss", "13:37:38"),
    ("time am/pm", T, "h:mm AM/PM", "1:37 PM"),
    ("elapsed hours", 1.5, "[h]:mm", "36:00"),
    ("datetime", D + T, "yyyy-mm-dd h:mm", "2026-08-04 13:37"),
    ("text placeholder", "hello", "@", "hello"),
]
for note, value, fmt, shows in cases:
    case(ws, row, note, value, {"numfmt": fmt, "shows": shows},
         number_format=fmt)
    row += 1
# Formats asserted only for survival of the format STRING (rendering is
# column-width/fill dependent or engine-specific).
for note, value, fmt in [
    ("accounting (fill char — no shows check)", 1234.5,
     '_-* #,##0.00_-;-* #,##0.00_-;_-* "-"??_-;_-@_-'),
    ("locale-prefixed date (no shows check)", D, "[$-409]d-mmm-yy"),
]:
    case(ws, row, note, value, {"numfmt": fmt}, number_format=fmt)
    row += 1

# ---------------------------------------------------------------------------
# Alignment
# ---------------------------------------------------------------------------
ws = sheet("Alignment")
row = 2
aligns = [
    ("left", {"halign": "left"}, Alignment(horizontal="left")),
    ("center", {"halign": "center"}, Alignment(horizontal="center")),
    ("right", {"halign": "right"}, Alignment(horizontal="right")),
    ("justify", {"halign": "justify"}, Alignment(horizontal="justify")),
    ("fill", {"halign": "fill"}, Alignment(horizontal="fill")),
    ("center continuous", {"halign": "centercontinuous"},
     Alignment(horizontal="centerContinuous")),
    ("distributed", {"halign": "distributed"},
     Alignment(horizontal="distributed")),
    ("top", {"valign": "top"}, Alignment(vertical="top")),
    ("middle", {"valign": "center"}, Alignment(vertical="center")),
    ("bottom (explicit)", {"valign": "bottom"}, Alignment(vertical="bottom")),
    ("v justify", {"valign": "justify"}, Alignment(vertical="justify")),
    ("wrap text", {"wrap": True},
     Alignment(wrap_text=True)),
    ("center + middle + wrap",
     {"halign": "center", "valign": "center", "wrap": True},
     Alignment(horizontal="center", vertical="center", wrap_text=True)),
]
for note, spec, aobj in aligns:
    value = ("wrapped text long enough to need more than one line"
             if spec.get("wrap") else note)
    case(ws, row, note, value, spec, style={"alignment": aobj})
    row += 1

# ---------------------------------------------------------------------------
# Borders
# ---------------------------------------------------------------------------
ws = sheet("Borders")
row = 2


def border_case(note, spec_border, border_obj):
    global row
    case(ws, row, note, note, {"border": spec_border},
         style={"border": border_obj})
    row += 2  # blank row between border cases so edges stay readable


for edge in ["left", "right", "top", "bottom"]:
    border_case(f"thin {edge} only", {edge: {"style": "thin"}},
                Border(**{edge: Side(style="thin")}))
border_case("thin box",
            {e: {"style": "thin"} for e in ["left", "right", "top", "bottom"]},
            Border(left=Side(style="thin"), right=Side(style="thin"),
                   top=Side(style="thin"), bottom=Side(style="thin")))
for bstyle in ["medium", "thick", "double", "dotted", "mediumDashed"]:
    border_case(f"{bstyle} box",
                {e: {"style": bstyle.lower()}
                 for e in ["left", "right", "top", "bottom"]},
                Border(left=Side(style=bstyle), right=Side(style=bstyle),
                       top=Side(style=bstyle), bottom=Side(style=bstyle)))
border_case("red thin box",
            {e: {"style": "thin", "color": "#FF0000"}
             for e in ["left", "right", "top", "bottom"]},
            Border(left=Side(style="thin", color="FFFF0000"),
                   right=Side(style="thin", color="FFFF0000"),
                   top=Side(style="thin", color="FFFF0000"),
                   bottom=Side(style="thin", color="FFFF0000")))
border_case("mixed edges (thick top, dotted bottom)",
            {"top": {"style": "thick"}, "bottom": {"style": "dotted"}},
            Border(top=Side(style="thick"), bottom=Side(style="dotted")))

# ---------------------------------------------------------------------------
# Merges
# ---------------------------------------------------------------------------
ws = sheet("Merges")
merges = [
    ("2x1 horizontal", "C3:D3"),
    ("1x3 vertical", "C5:C7"),
    ("3x3 block", "C9:E11"),
    ("wide banner", "C13:H13"),
]
row = 3
specs = [(3, merges[0]), (5, merges[1]), (9, merges[2]), (13, merges[3])]
for r, (note, rng) in specs:
    ws.cell(row=r, column=1, value=f"MER{r:03d}")
    ws.cell(row=r, column=2, value=note)
    anchor = rng.split(":")[0]
    ws[anchor] = note
    ws.merge_cells(rng)
    ws.cell(row=r, column=10, value=json.dumps({"merge": rng}))
ws["J1"] = "Spec"
# styled merge: banner with fill + centered
ws["C15"] = "styled merged banner"
ws.merge_cells("C15:F16")
ws["C15"].fill = PatternFill(fill_type="solid", fgColor="FFBDD7EE")
ws["C15"].alignment = Alignment(horizontal="center", vertical="center")
ws.cell(row=15, column=1, value="MER015")
ws.cell(row=15, column=2, value="merged + fill + centered")
ws.cell(row=15, column=10, value=json.dumps(
    {"merge": "C15:F16", "fill": "#BDD7EE", "halign": "center",
     "valign": "center"}))

# ---------------------------------------------------------------------------
# Geometry — column widths, row heights, hidden rows/cols
# ---------------------------------------------------------------------------
ws = sheet("Geometry")
row = 2
for col_letter, col_idx, chars in [("F", 6, 5), ("G", 7, 20), ("H", 8, 40)]:
    ws.column_dimensions[col_letter].width = chars
    ws.cell(row=row, column=col_idx, value=f"w={chars}")
    case(ws, row, f"column {col_letter} width {chars} chars", "",
         {"colwidth": {"col": col_idx, "chars": chars}})
    row += 1
for pt in [10, 30, 60]:
    ws.row_dimensions[row].height = pt
    case(ws, row, f"row height {pt}pt", f"{pt}pt tall",
         {"rowheight_pt": pt})
    row += 1
ws.row_dimensions[row].hidden = True
case(ws, row, "hidden row", "you should not see this",
     {"rowhidden": True})
row += 1
ws.column_dimensions["I"].hidden = True
ws.cell(row=row, column=9, value="hidden col")
case(ws, row, "hidden column I", "", {"colhidden": 9})
row += 1

# ---------------------------------------------------------------------------
# Frozen panes — its own sheet so the pane split is the only feature
# ---------------------------------------------------------------------------
ws = sheet("Frozen")
for r in range(1, 12):
    for c in range(1, 8):
        ws.cell(row=r, column=c, value=f"r{r}c{c}")
ws.freeze_panes = "C3"  # 2 frozen rows, 2 frozen columns
ws["J1"] = "Spec"
ws["J2"] = json.dumps({"frozen": {"rows": 2, "cols": 2}})

# ---------------------------------------------------------------------------
# Gaps — features IronCalc doesn't model. Only assert the VALUE survives
# the load (and the round-trip); the styling itself is expected to drop.
# ---------------------------------------------------------------------------
ws = sheet("Gaps")
row = 2
gaps = [
    ("text rotation 45°", {"alignment": Alignment(text_rotation=45)}),
    ("indent 3", {"alignment": Alignment(horizontal="left", indent=3)}),
    ("superscript", {"font": Font(vertAlign="superscript")}),
    ("pattern fill gray125",
     {"fill": PatternFill(fill_type="gray125", fgColor="FF808080")}),
    ("gradient fill",
     {"fill": GradientFill(stop=["FFFFFFFF", "FF1F4E79"])}),
    ("diagonal border",
     {"border": Border(diagonal=Side(style="thin"), diagonalUp=True)}),
    ("shrink to fit", {"alignment": Alignment(shrink_to_fit=True)}),
]
for note, style in gaps:
    case(ws, row, note, f"survives: {note}", {"survive": True}, style=style)
    row += 1

# ---------------------------------------------------------------------------

del wb["Sheet"]
out = OUT / "style_gauntlet.xlsx"
wb.save(out)
n_sheets = len(wb.sheetnames)
print(f"wrote {out} — {n_sheets} sheets: {', '.join(wb.sheetnames)}")
